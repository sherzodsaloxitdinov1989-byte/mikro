#!/usr/bin/env python3
"""«Mikroiqtisodiyot» darsligini (DOCX) platforma ma'lumotlariga o'giradi.

Darslikning har bir bobida tayyor test, glossariy, masalalar, mustaqil ta'lim va nazorat savollari bor.
Skript ularni DARSLIKNING O'ZIDAN ajratib oladi (hech narsa o'ylab topilmaydi):

    python3 scripts/import-mikro.py manbalar/darslik.docx

Yoziladi:
    public/data/book.js          — bob matni (bo'limlar, rasmlar, jadvallar, «keys» qutilari)
    public/data/layers-auto.js   — glossariy, testlar (kalitlar bilan), metodlar, mustaqil ishlar
    public/data/practice-auto.js — masalalar (namunaviy yechimlari bilan)
    public/img/book/*.png        — darslikdagi grafiklar
"""
import html as H
import json
import re
import sys
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "public" / "data"
IMG = ROOT / "public" / "img" / "book"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

CH_RE = re.compile(r"^(?:[IVX]+|\d+)\s*-?\s*BOB\.?\s*(.*)$", re.I)
SEC_RE = re.compile(r"^\d{1,2}\.\d{1,2}\.?\s")
esc = lambda s: H.escape(s, quote=False)


def clean(s):
    return re.sub(r"\s+", " ", s).strip()


class Doc:
    def __init__(self, path):
        self.d = docx.Document(str(path))
        self.rels = self.d.part.rels
        IMG.mkdir(parents=True, exist_ok=True)
        self.saved = set()

    def run_html(self, p_el):
        """Xatboshidan matn (yuqori/pastki indekslar bilan) va rasm nomlari."""
        out, imgs = [], []
        for r in p_el.iter(W + "r"):
            va = r.find(f"{W}rPr/{W}vertAlign")
            mode = va.get(W + "val") if va is not None else None
            txt = ""
            for ch in r:
                if ch.tag == W + "t":
                    txt += ch.text or ""
                elif ch.tag == W + "tab":
                    txt += " "
                elif ch.tag == W + "br":
                    txt += "\n"
            if txt:
                e = esc(txt).replace("\n", "<br>")
                out.append(f"<sup>{e}</sup>" if mode == "superscript" else f"<sub>{e}</sub>" if mode == "subscript" else e)
            for bl in r.iter(A_BLIP):
                rid = bl.get(R_EMBED)
                if rid and rid in self.rels:
                    part = self.rels[rid].target_part
                    name = Path(str(part.partname)).name
                    if name not in self.saved:
                        (IMG / name).write_bytes(part.blob)
                        self.saved.add(name)
                    imgs.append(name)
        return "".join(out), imgs

    def items(self):
        """Hujjat tanasi ketma-ketligi: h1/h2/h3/p/tbl."""
        for ch in self.d.element.body.iterchildren():
            tag = ch.tag.split("}")[-1]
            if tag == "p":
                p = docx.text.paragraph.Paragraph(ch, self.d)
                st = (p.style.name or "").lower() if p.style is not None else ""
                kind = "h1" if "heading 1" in st else "h2" if "heading 2" in st else "h3" if "heading 3" in st else "p"
                html, imgs = self.run_html(ch)
                yield {"k": kind, "text": clean(p.text), "html": html.strip(), "imgs": imgs}
            elif tag == "tbl":
                yield {"k": "tbl", "tbl": docx.table.Table(ch, self.d)}


def cell_paras(cell):
    return [clean(p.text) for p in cell.paragraphs if clean(p.text)]


def table_html(doc, tb):
    """Bitta kataklik quti → «callout»; ko'p kataklik → jadval."""
    if len(tb.columns) == 1 and len(tb.rows) == 1:
        c = tb.rows[0].cells[0]
        parts = []
        for p in c.paragraphs:
            h, imgs = doc.run_html(p._p)
            if h.strip():
                parts.append(h.strip())
            for n in imgs:
                parts.append(f'<img src="/img/book/{n}" loading="lazy" alt="">')
        if not parts:
            return ""
        title, rest = parts[0], parts[1:]
        cls = "callout"
        if clean(c.paragraphs[0].text).upper().startswith("RAQAMLAR TILIDA"):
            cls += " callout-calc"
        body, bullets = "", []
        for x in rest:
            if x.startswith("•"):
                bullets.append(f"<li>{x.lstrip('• ').strip()}</li>")
                continue
            if bullets:
                body += "<ul>" + "".join(bullets) + "</ul>"
                bullets = []
            body += x if x.startswith("<img") else f"<p>{x}</p>"
        if bullets:
            body += "<ul>" + "".join(bullets) + "</ul>"
        return f'<div class="{cls}"><p><b>{title}</b></p>{body}</div>'
    rows = []
    for r in tb.rows:
        cells, last = [], None
        for c in r.cells:
            if c._tc is last:
                continue
            last = c._tc
            cells.append(clean(c.text))
        rows.append(cells)
    if not rows:
        return ""
    head = "".join(f"<th>{esc(x)}</th>" for x in rows[0])
    bodyr = "".join("<tr>" + "".join(f"<td>{esc(x)}</td>" for x in r) + "</tr>" for r in rows[1:])
    return f'<div class="tbl-wrap"><table><thead><tr>{head}</tr></thead><tbody>{bodyr}</tbody></table></div>'


class Render:
    def __init__(self, doc):
        self.doc, self.out, self.bul = doc, [], []

    def flush(self):
        if self.bul:
            self.out.append("<ul>" + "".join(self.bul) + "</ul>")
            self.bul = []

    def add(self, it):
        k = it["k"]
        if k == "tbl":
            self.flush()
            h = table_html(self.doc, it["tbl"])
            if h:
                self.out.append(h)
            return
        for n in it.get("imgs", []):
            self.flush()
            self.out.append(f'<p class="figure"><img src="/img/book/{n}" loading="lazy" alt=""></p>')
        t, h = it["text"], it["html"]
        if not t:
            return
        if k in ("h2", "h3"):
            self.flush()
            self.out.append(f"<h4>{esc(t)}</h4>")
        elif t.startswith("•"):
            self.bul.append(f"<li>{h.lstrip('• ').strip()}</li>")
        else:
            self.flush()
            if re.match(r"^\d{1,2}(\.\d{1,2})*\.?\s*-?\s*(jadval|rasm|chizma|diagramma)\b", t, re.I):
                self.out.append(f'<p class="caption"><b>{h}</b></p>')
            elif re.match(r"^(Manba|Izoh)\s*:", t):
                self.out.append(f'<p class="source">{h}</p>')
            else:
                self.out.append(f"<p>{h}</p>")

    def html(self):
        self.flush()
        return "".join(self.out)


def numbered(items):
    """«1. matn» ko'rinishidagi xatboshilar → [[raqam, matn]]"""
    out = []
    for it in items:
        if it["k"] != "p" or not it["text"]:
            continue
        m = re.match(r"^(\d{1,2})[.)]\s*(.+)$", it["text"])
        if m:
            out.append([int(m.group(1)), m.group(2)])
        elif out:
            out[-1][1] += " " + it["text"]
    return out


def sentence_case(s):
    s = s.lower()
    s = re.sub(r"([.:]\s+)([a-zо‘ʻ'’])", lambda m: m.group(1) + m.group(2).upper(), s)
    for i, ch in enumerate(s):
        if ch.isalpha():
            return s[:i] + ch.upper() + s[i + 1:]
    return s


def main():
    path = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT / "manbalar" / "darslik.docx")
    doc = Doc(path)
    items = list(doc.items())

    intro_items, chapters, keys_items, lit_items = [], [], [], []
    cur, zone = None, "pre"
    for it in items:
        if it["k"] == "h1":
            t = it["text"]
            m = CH_RE.match(t)
            if m:
                cur = {"title": clean(m.group(1)).strip(" .:"), "items": []}
                chapters.append(cur)
                zone = "ch"
            elif t.upper().startswith("KIRISH"):
                zone = "intro"
            elif "KALIT" in t.upper() or "JAVOBLARI" in t.upper():
                zone = "keys"
            elif "ADABIYOT" in t.upper():
                zone = "lit"
            else:
                zone = "skip"
            continue
        if zone == "ch":
            cur["items"].append(it)
        elif zone == "intro":
            intro_items.append(it)
        elif zone == "keys":
            keys_items.append(it)
        elif zone == "lit":
            lit_items.append(it)

    # Test kalitlari
    keys, ch_no = {}, None
    for it in keys_items:
        if it["k"] == "p":
            m = re.match(r"^(\d+)-bob", it["text"])
            if m:
                ch_no = int(m.group(1))
        elif it["k"] == "tbl" and ch_no:
            rows = [[clean(c.text) for c in r.cells] for r in it["tbl"].rows]
            d = {}
            for i in range(0, len(rows) - 1, 2):
                for n, a in zip(rows[i], rows[i + 1]):
                    if n.isdigit() and a in ("A", "B", "C", "D"):
                        d[int(n)] = "ABCD".index(a)
            keys[ch_no] = d

    topics, layers, practice = [], {}, {}
    for num, ch in enumerate(chapters, 1):
        parts, name, buf, order = {}, "_pre", [], []
        for it in ch["items"]:
            if it["k"] == "h2":
                parts.setdefault(name, []).extend(buf)
                name, buf = it["text"], []
                order.append(name)
            else:
                buf.append(it)
        parts.setdefault(name, []).extend(buf)

        def find(*ks):
            for n in parts:
                if any(k in n.lower() for k in ks):
                    return parts[n]
            return []

        sections, plan = [], []
        pre = Render(doc)
        for it in parts.get("_pre", []):
            pre.add(it)
        if pre.html():
            sections.append({"title": "", "html": pre.html()})
        for n in order:
            if SEC_RE.match(n) or "xulosalar" in n.lower():
                r = Render(doc)
                for it in parts[n]:
                    r.add(it)
                sections.append({"title": n.rstrip("."), "html": r.html()})
                if SEC_RE.match(n):
                    plan.append(n.rstrip("."))
        qs = [clean(q) for _, q in numbered(find("nazorat savollari"))]

        gloss = []
        for it in find("glossariy"):
            if it["k"] == "p" and it["text"]:
                a = re.split(r"\s[–—-]\s", it["text"], maxsplit=1)
                if len(a) == 2:
                    gloss.append((clean(a[0]), clean(a[1])))

        raw = []
        for it in find("test savollari"):
            if it["k"] != "p" or not it["text"]:
                continue
            t = it["text"]
            mo = re.match(r"^([ABCD])\)\s*(.+)$", t)
            mq = re.match(r"^(\d{1,2})[.)]\s*(.+)$", t)
            if mo and raw:
                raw[-1]["o"].append(mo.group(2))
            elif mq:
                raw.append({"n": int(mq.group(1)), "q": mq.group(2), "o": []})
            elif raw and not raw[-1]["o"]:
                raw[-1]["q"] += " " + t
        tests = []
        for r in raw:
            k = keys.get(num, {}).get(r["n"])
            if k is not None and len(r["o"]) == 4:
                tests.append({"q": r["q"], "options": r["o"], "c": k})

        selfs = []
        for i, (_, t) in enumerate(numbered(find("mustaqil")), 1):
            m = re.match(r"^[«“\"]([^»”\"]{5,90})[»”\"]", t)
            title = m.group(1) if m else clean(re.split(r"[.:;]", t)[0])[:80]
            selfs.append({"id": f"t{num}-s{i}", "title": title, "type": "Mustaqil ish", "description": t})

        probs = numbered(find("masala va mashq"))
        sols = {}
        for it in find("masalalar yechimi"):
            if it["k"] == "tbl" and len(it["tbl"].columns) == 1:
                ps = cell_paras(it["tbl"].rows[0].cells[0])
                m = re.search(r"(\d+)-masala", ps[0]) if ps else None
                if m:
                    sols[int(m.group(1))] = ps
        concl = [re.sub(r"^\d+\)\s*", "", it["text"]) for it in find("xulosalar") if it["k"] == "p" and it["text"]]
        boxes = []
        for it in ch["items"]:
            if it["k"] == "tbl" and len(it["tbl"].columns) == 1:
                ps = cell_paras(it["tbl"].rows[0].cells[0])
                if ps and re.match(r"^O[‘'’`ʻ]ZBEKISTON MISOLIDA", ps[0].upper()):
                    boxes.append(ps)

        title = sentence_case(ch["title"])
        methods = []
        sample = [re.sub(r"^\d+\.\d+\.?\s*", "", p) for p in plan][:8]
        if sample:
            methods.append({"type": "cluster", "id": f"t{num}-cluster", "title": f"Klaster: {title}", "instruction": "Mavzuning asosiy tushunchalari va ularning o'zaro bog'liqligini klaster ko'rinishida tuzing.", "center": title[:40], "sample": sample})
        pairs = [[t, d.rstrip(".") + "."] for t, d in gloss if len(d) <= 160][:6]
        if len(pairs) >= 4:
            methods.append({"type": "matching", "id": f"t{num}-match", "title": "Tushuncha va uning ta'rifi", "instruction": "Har bir tushunchani uning ta'rifi bilan moslang.", "pairs": pairs})
        if probs and 1 in sols and len(sols[1]) > 1:
            methods.append({"type": "case", "id": f"t{num}-case", "title": "Keys: amaliy masala", "instruction": "Masala shartini o'qing, yechim yo'lini o'zingiz yozing, so'ng namunaviy yechim bilan solishtiring.", "text": probs[0][1], "questions": ["Masalaning har bir bandini ketma-ket yeching va izohlang.", "Olingan natijadan qanday iqtisodiy xulosa kelib chiqadi?"], "model": " ".join(sols[1][1:])})
        elif boxes and concl:
            b = boxes[0]
            methods.append({"type": "case", "id": f"t{num}-case", "title": "Keys: O'zbekiston misoli", "instruction": "Vaziyatni o'qing va mavzudagi tushunchalar asosida tahlil qiling.", "text": " ".join(b[1:])[:1400], "questions": ["Bu vaziyatni mavzudagi qaysi tushunchalar yordamida izohlash mumkin?", "Qanday iqtisodiy xulosa va tavsiya berasiz?"], "model": concl[0]})
        sents = [c for c in concl if 40 < len(c) < 260][:5]
        if len(sents) >= 3:
            methods.append({"type": "insert", "id": f"t{num}-insert", "title": "INSERT: bob xulosalari", "instruction": "Har bir fikrni o'qib, belgilang: ✓ bilaman, + yangi, − qarama-qarshi, ? tushunarsiz.", "sentences": sents})

        topics.append({"num": num, "chapter": "", "title": ch["title"], "plan": plan, "sections": sections or [{"title": "", "html": ""}], "questions": qs, "keywords": [g[0] for g in gloss][:14]})
        layers[num] = {"icon": "", "title": title, "glossary": gloss, "methods": methods, "tests": tests, "selfStudy": selfs, "resources": []}
        if probs or qs:
            open_ = [t for _, t in probs]
            refs = [" ".join(sols[n][1:]) if sols.get(n) and len(sols[n]) > 1 else "" for n, _ in probs]
            if not open_:
                open_ = qs[:3]
                refs = [""] * len(open_)
            practice[num] = {
                "goal": f"«{title}» mavzusi bo'yicha masala va topshiriqlarni yechib, nazariy bilimlarni amalda qo'llash.",
                "concepts": [g[0] for g in gloss][:6],
                "source": f"Saloxitdinov Sh.F. Mikroiqtisodiyot, {num}-bob: «Masala va mashqlar»",
                "tasks": [],
                "open": open_,
                "refs": refs,
            }

    r = Render(doc)
    for it in intro_items:
        r.add(it)
    lit = [it["text"] for it in lit_items if it["k"] == "p" and it["text"]]

    js = lambda o: json.dumps(o, ensure_ascii=False, indent=1)
    (DATA / "book.js").write_text(
        "// Avtomatik yaratilgan fayl — qo'lda tahrirlamang.\n"
        f"// Manba: {path.name}\n// Qayta yaratish: python3 scripts/import-mikro.py manbalar/darslik.docx\n"
        "export const BOOK = " + js({"intro": r.html(), "topics": topics, "literature": lit}) + ";\n",
        encoding="utf-8",
    )
    (DATA / "layers-auto.js").write_text(
        "// Avtomatik yaratilgan (scripts/import-mikro.py): darslikdagi glossariy, testlar (kalitlar bilan), metodlar, mustaqil ishlar.\n"
        "export const AUTO = " + js(layers) + ";\n",
        encoding="utf-8",
    )
    (DATA / "practice-auto.js").write_text(
        "// Avtomatik yaratilgan (scripts/import-mikro.py): darslikdagi «Masala va mashqlar» va ularning yechimlari.\n"
        "export const AUTO_PRACTICE = " + js(practice) + ";\n",
        encoding="utf-8",
    )
    print(f"{'№':>3} {'bo‘lim':>6} {'savol':>5} {'test':>4} {'gloss':>5} {'masala':>6} {'mustaqil':>8} {'metod':>5}  nomi")
    for t in topics:
        L, P = layers[t["num"]], practice.get(t["num"], {})
        print(f"{t['num']:>3} {len(t['sections']):>6} {len(t['questions']):>5} {len(L['tests']):>4} {len(L['glossary']):>5} {len(P.get('open', [])):>6} {len(L['selfStudy']):>8} {len(L['methods']):>5}  {t['title'][:55]}")
    print(f"rasmlar: {len(doc.saved)} ta; adabiyotlar: {len(lit)}; kirish: {len(r.html())} belgi")


if __name__ == "__main__":
    main()
